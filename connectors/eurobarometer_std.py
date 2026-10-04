"""Standard Eurobarometer, Volume A (European Commission). Licence: Commission reuse decision 2011/833/EU, CC BY 4.0
("© European Union, Standard Eurobarometer"). One object per wave × country with the measured indicators, the
fieldwork dates and the sample: trust in the national government and parliament, things going in the right/wrong
direction, the current national economy, expectations for the next 12 months, the two most important issues.
Indicators are recognised by the ORIGINAL QUESTION TEXT (question numbers change between waves). Values are the
published weighted shares; nothing is modelled or inferred."""

import io
import re
import zipfile

import openpyxl

from connectors.obs_common import NAMES_IT, iso, record
from nexum.core.scheduler import FetchRequest

VERSION = "1.0.0"
# Volume A of the last four standard waves (data.europa.eu datasets s3215_102_2, s3372_103_3, s3378_104_1, s3613_105_2)
WAVES = [("102.2", "A2534D981B6C0CF8878DEB8447FB08F4"), ("103.3", "AF452C0A6DD70E5D13273C5DF2086956"),
         ("104.1", "A68C45C80315C6F37BBA366C32491DED"), ("105.2", "99134D6EF7DB334B8485DC4C0535E9C2")]
URL = "https://webgate.ec.europa.eu/ebsm/api/public/odp/download?key={}"
SKIP = {"EU27", "UE27", "DEW", "DEE", "CY_TCC", "CY(TCC)", "XK"}
# (code, Italian label, question pattern, sub-item pattern, answer rows kept)
TOPIC = {"trust": "istituzioni", "direction": "direzione", "economy": "economia", "life": "futuro personale", "democracy": "istituzioni"}
SPEC = [
    ("trust.government", "Fiducia nel governo nazionale", r"how much trust do you have", r"^The \(NATIONALITY\) Government$", {"Tend to trust": "tende a fidarsi", "Tend not to trust": "tende a non fidarsi"}),
    ("trust.parliament", "Fiducia nel parlamento nazionale", r"how much trust do you have", r"^The \(NATIONALITY ?PARLIAMENT\)$|^The \(NATIONALITY\) Parliament$", {"Tend to trust": "tende a fidarsi", "Tend not to trust": "tende a non fidarsi"}),
    ("direction.country", "Direzione del Paese", r"right direction or in the wrong direction", r"^\(OUR COUNTRY\)$", {"Things are going in the right direction": "direzione giusta", "Things are going in the wrong direction": "direzione sbagliata"}),
    ("economy.now", "Situazione attuale dell'economia nazionale", r"judge the current situation", r"situation of the \(NATIONALITY\) economy", {"Total 'Good'": "buona", "Total 'Bad'": "cattiva"}),
    ("economy.expect12", "Economia nazionale nei prossimi 12 mesi", r"expectations for the next twelve months", r"^The economic situation in \(OUR COUNTRY\)$", {"Better": "migliore", "Worse": "peggiore", "Same": "uguale"}),
    ("life.expect12", "La propria vita nei prossimi 12 mesi", r"expectations for the next twelve months", r"^Your life in general$", {"Better": "migliore", "Worse": "peggiore", "Same": "uguale"}),
    ("democracy.satisfaction", "Soddisfazione per il funzionamento della democrazia", r"satisfied.*the way democracy works in \(OUR COUNTRY\)", r".*", {"Total 'Satisfied'": "soddisfatti", "Total 'Not satisfied'": "non soddisfatti"}),
]
ISSUES = r"two most important issues facing \(OUR COUNTRY\)"
ISSUE_IT = {"crime": "criminalità", "the economic situation": "situazione economica", "rising prices/inflation/cost of living": "aumento dei prezzi, inflazione, costo della vita",
            "taxation": "tasse", "unemployment": "disoccupazione", "terrorism": "terrorismo", "housing": "abitazione",
            "government debt": "debito pubblico", "immigration": "immigrazione", "health": "salute", "the education system": "istruzione",
            "pensions": "pensioni", "the environment, climate and energy issues": "ambiente, clima ed energia",
            "energy supply": "approvvigionamento energetico", "the international situation": "situazione internazionale",
            "security and defence": "sicurezza e difesa", "the environment and climate change": "ambiente e cambiamento climatico",
            "the environment and risks related to climate change": "ambiente e rischi legati al cambiamento climatico",
            "threats to democracy (for example through information manipulation)": "minacce alla democrazia (per esempio manipolazione dell'informazione)",
            "russia’s invasion of ukraine": "invasione russa dell'Ucraina", "russia's invasion of ukraine": "invasione russa dell'Ucraina",
            "the conflict in the middle east": "conflitto in Medio Oriente", "other": "altro", "none": "nessuno", "don't know": "non sa"}
ISSUE_IT_BY_SLUG = {re.sub(r"[^a-z0-9]+", "_", k).strip("_")[:40]: v for k, v in ISSUE_IT.items()}


def describe():
    return {"connector_version": VERSION, "produces": ["observation.survey_wave"]}


def plan(mode, state, source, today=None):
    return [FetchRequest(URL.format(k), f"eb_{w}_volume_A.xlsx") for w, k in WAVES]


def next_state(state, request, result, today=None):
    return state


def _slug(s):
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")[:40]


def _dates(fw):
    """'Terrain/Fieldwork : 12/3 - 5/4/2026' → ('2026-03-12', '2026-04-05') (start month may be in the previous year)."""
    m = re.search(r"(\d{1,2})/(\d{1,2})(?:/(\d{4}))?\s*-\s*(\d{1,2})/(\d{1,2})/(\d{4})", fw or "")
    if not m:
        return None, None
    d1, m1, y1, d2, m2, y2 = m.groups()
    y2 = int(y2)
    y1 = int(y1) if y1 else (y2 - 1 if int(m1) > int(m2) else y2)
    return f"{y1:04d}-{int(m1):02d}-{int(d1):02d}", f"{y2:04d}-{int(m2):02d}-{int(d2):02d}"


def _table(rows, wave):
    """Country → column of the CURRENT wave, the row of the totals, the English answer rows."""
    ci = next((i for i, r in enumerate(rows[:16]) if any(str(x or "").strip() == "BE" for x in r)), None)
    if ci is None:
        return None
    codes, nxt = rows[ci], rows[ci + 1] if ci + 1 < len(rows) else ()
    cols = {}
    if any(isinstance(x, str) and x.startswith("EB - ") for x in nxt):        # two columns per country (old/new wave)
        tag = f"EB - {wave}"
        last = None
        for j, x in enumerate(codes):
            if x:
                last = str(x).strip()
            if j < len(nxt) and nxt[j] == tag and last:
                cols[last.split()[-1]] = j
        start = ci + 2
    else:
        for j, x in enumerate(codes):
            if j >= 2 and x:
                cols[str(x).strip().split()[-1]] = j
        start = ci + 1
    k = next((i for i in range(start - 1, min(len(rows), start + 4)) if rows[i] and rows[i][1] == "Total"), None)
    if k is None:
        return None
    return cols, k, rows


def _val(row, j):
    v = row[j] if j < len(row) else None
    return float(v) if isinstance(v, (int, float)) and 0 <= v <= 1 else None


def parse(data: bytes, meta: dict):
    if data[:2] == b"PK" and b"xl/workbook" not in data[:4096] and not data[:4096].count(b"[Content_Types]"):
        try:
            z = zipfile.ZipFile(io.BytesIO(data))
            names = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
            if names and "xl/workbook.xml" not in z.namelist():
                data = z.read(names[0])
        except zipfile.BadZipFile:
            return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    by_country, fieldwork, wave_label = {}, (None, None), None
    for sh in wb.sheetnames[1:]:
        rows = [tuple(r) for r in wb[sh].iter_rows(values_only=True)]
        if len(rows) < 12:
            continue
        top = " ".join(str(x) for x in rows[0] if x)
        wm = re.search(r"Eurobarometer\s*-\s*([\d.]+)", top)
        if not wm:
            continue
        wave = wm.group(1)
        wave_label = wave
        if fieldwork == (None, None):
            fieldwork = _dates(" ".join(str(x) for x in rows[1] if x))
        q, item = str(rows[2][7] or ""), str(rows[3][7] or "").strip()
        specs = [s for s in SPEC if re.search(s[2], q, re.I) and re.search(s[3], item)]
        is_issues = re.search(ISSUES, q, re.I)
        if not specs and not is_issues:
            continue
        t = _table(rows, wave)
        if not t:
            continue
        cols, k, rows = t
        english = [rows[i] for i in range(k + 3, len(rows), 2) if rows[i] and isinstance(rows[i][1], str)]
        for code_raw, j in cols.items():
            if code_raw.upper() in SKIP:
                continue
            c = iso(code_raw)
            if not c:
                continue
            n = rows[k][j] if isinstance(rows[k][j], (int, float)) else None
            bag = by_country.setdefault(c, {"n": n, "items": []})
            if n and not bag["n"]:
                bag["n"] = n
            if specs:
                code, label, _, _, answers = specs[0]
                for r in english:
                    a = r[1].strip()
                    if a in answers and _val(r, j) is not None:
                        bag["items"].append([f"{code}:{_slug(a)}", f"{label}: {answers[a]}", f"{q} — {item}" if item else q, a,
                                             round(_val(r, j) * 100, 1), n, TOPIC[code.split(".")[0]]])
            else:
                for r in english:
                    a = r[1].strip()
                    v = _val(r, j)
                    if v is None or a.lower() in ("other (spontaneous)", "none (spontaneous)", "don't know", "total"):
                        continue
                    it = ISSUE_IT_BY_SLUG.get(_slug(a), a)
                    bag["items"].append([f"concern:{_slug(a)}", f"Problema principale: {it}", q, a, round(v * 100, 1), n, "preoccupazioni"])
    if not wave_label:
        return
    season = "primavera" if fieldwork[1] and fieldwork[1][5:7] <= "07" else "autunno"
    for c, bag in sorted(by_country.items()):
        if not bag["items"]:
            continue
        name = NAMES_IT.get(c, c)
        yield record(meta, f"eb:{wave_label}:{c}", "observation.survey_wave",
                     f"Eurobarometro standard {wave_label} ({season} {fieldwork[1][:4] if fieldwork[1] else ''}) · {name}", c,
                     {"dataset": f"Eurobarometro standard {wave_label}", "wave": wave_label,
                      "fieldwork_start": fieldwork[0], "fieldwork_end": fieldwork[1], "sample_size": bag["n"],
                      "population": "residenti di 15 anni e più", "method": "interviste faccia a faccia (in alcuni Paesi anche online), risultati ponderati",
                      "probability_sample": True, "items": bag["items"]}, scheme="obs_wave",
                     text=f"Eurobarometro {name} fiducia governo direzione economia problemi")
