"""European Commission (DG ENER) Weekly Oil Bulletin — consumer prices of Euro-super 95 and automotive gas oil in the
27 EU member states, weekly, inclusive of duties and taxes (price history file), with the VAT rate and the excise
duty in force as published (each with its "since" date). Licence: CC BY 4.0 (Commission Decision 2011/833/EU);
"Source: European Commission, Weekly Oil Bulletin". Prices in the history file are in EUR per 1 000 litres: divided
by 1 000 (EUR per litre), nothing else. Excise duties are in national currency per 1 000 litres, as published."""

import datetime
import io
import re

import openpyxl

from connectors.fuel_common import fuel
from connectors.obs_common import iso
from nexum.core.scheduler import FetchRequest

VERSION = "1.0.0"
PAGE = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
FUELS = {"euro95": ("petrol", 0), "diesel": ("diesel", 1)}            # column suffix → fuel, tax column offset
# national currencies of the excise tables (euro area members publish in EUR); Bulgaria adopted the euro on 2026-01-01
EU27 = set("AT BE BG CY CZ DE DK EE ES FI FR GR HR HU IE IT LT LU LV MT NL PL PT RO SE SI SK".split())
CURRENCY = {"CZ": "CZK", "DK": "DKK", "HU": "HUF", "PL": "PLN", "RO": "RON", "SE": "SEK"}


def describe():
    return {"connector_version": VERSION, "produces": ["observation.fuel_price"]}


def plan(mode, state, source, today=None):
    reqs = [FetchRequest(PAGE, "wob_page.html")]
    if (state or {}).get("history"):
        reqs.append(FetchRequest(state["history"], "wob_prices_history.xlsx"))
    return reqs


def page_info(data):
    m = re.search(rb'href="(/document/download/[0-9a-f-]+_en\?filename=Weekly_Oil_Bulletin_Prices_History[^"]*\.xlsx)"', data or b"")
    return {"history": "https://energy.ec.europa.eu" + m.group(1).decode()} if m else None


def next_state(state, request, result, today=None):
    st = dict(state or {})
    if result and result.get("history"):
        st["history"] = result["history"]
    return st


def _day(v):
    return v.date().isoformat() if isinstance(v, datetime.datetime) else None


def _latest_rates(ws):
    """{country: {column: (since, value)}} — per product, the latest row of the country block that has a value (a row
    lists only the products whose rate changed on that date; the newest rows come first)."""
    out, cur = {}, None
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0]:
            cur = iso(str(r[0]).rstrip("_"))
        if not cur or not isinstance(r[1], datetime.datetime):
            continue
        for k in (0, 1):
            v = r[2 + k] if len(r) > 2 + k else None
            if v is not None and k not in out.setdefault(cur, {}):
                out[cur][k] = (r[1].date().isoformat(), v)
    return out


def parse(data: bytes, meta: dict):
    if data[:2] != b"PK":
        return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = {}
    for name in ("Prices with taxes", "Prices wo taxes"):
        rows = list(wb[name].iter_rows(values_only=True))
        head = rows[0]
        cols = {}
        for j, h in enumerate(head):
            m = re.match(r"^([A-Z]{2})_price_(?:with|wo)_tax_(euro95|diesel)$", str(h or ""))
            if m and iso(m.group(1)) in EU27:          # never the EU aggregate, nor the UK (published until 2020)
                cols[(iso(m.group(1)), m.group(2))] = j
        series = {}
        for r in rows[3:]:
            d = _day(r[0])
            if not d:
                continue
            for k, j in cols.items():
                v = r[j] if j < len(r) else None
                if isinstance(v, (int, float)) and v > 0:
                    series.setdefault(k, []).append([d, d, v / 1000])
        sheets[name] = series
    vat, excise = _latest_rates(wb["VAT"]), _latest_rates(wb["Excise duties"])
    for (c, f), pts in sorted(sheets["Prices with taxes"].items()):
        fuel_type, k = FUELS[f]
        last = max(pts)[0]
        wo = {p[0]: p[2] for p in sheets["Prices wo taxes"].get((c, f), [])}
        taxes = []
        if k in vat.get(c, {}):
            since, v = vat[c][k]
            taxes.append(["IVA", float(v), "%", since])
        if k in excise.get(c, {}):
            since, v = excise[c][k]
            cur = "BGN" if c == "BG" and since < "2026-01-01" else CURRENCY.get(c, "EUR")
            taxes.append(["Accisa", round(float(v) / 1000, 4), f"{cur}/l", since])
        rec = fuel(meta, c, fuel_type, pts, unit="EUR/l", currency="EUR", frequency="settimanale",
                   definition="Prezzo medio al consumo con tasse e accise, rilevato il lunedì (Weekly Oil Bulletin); "
                              "pubblicato in EUR per 1 000 litri, diviso per 1 000",
                   dataset="Commissione europea, DG ENER — Weekly Oil Bulletin (storico dei prezzi)", taxes=taxes or None,
                   price_without_taxes=wo.get(last))
        if rec:
            yield rec
