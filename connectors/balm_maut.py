"""Germany — BALM Mauttabelle: every toll section of the truck toll on federal trunk roads (about 134 000 sections:
ID, from/to junction, length, road, federal state, coordinates), current version (the table changes about every six
weeks; the official page lists every version and is read first, backfill rounds).
Licence: Datenlizenz Deutschland – Namensnennung 2.0 (catalogue); "Bundesamt für Logistik und Mobilität (BALM)".

PACKAGING (O7): the sections are grouped per ROAD and FEDERAL STATE into one object each (a few thousand objects, not
134 000): its geometry is the sections' lines, its "segments" keep every section exactly (ID, from, to, km). A
section used in part is charged in full (BFStrMG § 3(2)); the amount of a section = its length × the rate of the
vehicle's class, rounded to the cent (§ 3(4)). NEXUM does not compute routes."""

import io
import re
import zipfile

import openpyxl

from connectors.base import content_version
from nexum.core.records import Assertion, NormalizedRecord, Target
from nexum.core.scheduler import FetchRequest

VERSION = "1.0.0"
PAGE = "https://www.balm.bund.de/DE/Themen/Lkw-Maut/Mauttabelle/mauttabelle_node.html"
LAND = {"BW": "Baden-Württemberg", "BY": "Baviera", "BE": "Berlino", "BB": "Brandeburgo", "HB": "Brema", "HH": "Amburgo",
        "HE": "Assia", "MV": "Meclemburgo-Pomerania Anteriore", "NI": "Bassa Sassonia", "NW": "Renania Settentrionale-Vestfalia",
        "RP": "Renania-Palatinato", "SL": "Saarland", "SN": "Sassonia", "ST": "Sassonia-Anhalt", "SH": "Schleswig-Holstein",
        "TH": "Turingia"}


def describe():
    return {"connector_version": VERSION, "produces": ["toll.road_section"]}


def plan(mode, state, source, today=None):
    reqs = [FetchRequest(PAGE, "balm_mauttabelle_page.html")]
    if (state or {}).get("table"):
        reqs.append(FetchRequest(state["table"], "balm_mauttabelle_current.zip"))
    return reqs


def page_info(data):
    links = re.findall(rb'href="(/SharedDocs/Downloads/DE/Lkw-Maut/Mauttabelle/Mauttabelle_(\d{4})[-_](\d{2})[-_](\d{2})\.zip[^"]*)"', data or b"")
    if not links:
        return None
    best = max(links, key=lambda m: (m[1], m[2], m[3]))                 # the latest version listed
    return {"table": "https://www.balm.bund.de" + best[0].decode().replace("&amp;", "&")}


def next_state(state, request, result, today=None):
    st = dict(state or {})
    if result and result.get("table"):
        st["table"] = result["table"]
    return st


def _chains(segs):
    """The sections' lines joined where one ends at the next one's start (fewer, longer lines; no point invented)."""
    starts = {}
    for s in segs:
        starts.setdefault(s[4], []).append(s)
    used, lines = set(), []
    ends = {s[5] for s in segs}
    order = [s for s in segs if s[4] not in ends] + segs                  # line heads first
    for s in order:
        if s[0] in used:
            continue
        line, cur = [s[4], s[5]], s
        used.add(s[0])
        while True:
            nxt = next((n for n in starts.get(cur[5], []) if n[0] not in used), None)
            if not nxt:
                break
            used.add(nxt[0])
            line.append(nxt[5])
            cur = nxt
        lines.append(line)
    return lines


def parse(data: bytes, meta: dict):
    if data[:2] != b"PK":
        return
    z = zipfile.ZipFile(io.BytesIO(data))
    wb = openpyxl.load_workbook(io.BytesIO(z.read(next(n for n in z.namelist() if n.endswith(".xlsx")))), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    title = str(next(rows)[0] or "")
    m = re.search(r"Version (\d+), gültig ab (\d{2})\.(\d{2})\.(\d{4})", title)
    version, valid_from = (m.group(1), f"{m.group(4)}-{m.group(3)}-{m.group(2)}") if m else (None, None)
    next(rows)
    groups = {}
    for r in rows:
        if not r or r[0] is None or r[4] is None:
            continue
        try:
            a = (round(float(r[7]), 5), round(float(r[6]), 5))
            b = (round(float(r[9]), 5), round(float(r[8]), 5))
            km = float(r[3])
        except (TypeError, ValueError):
            continue
        road, land = str(r[4]).strip(), str(r[5] or "").strip()
        groups.setdefault((road, land), []).append([int(r[0]), str(r[1] or ""), str(r[2] or ""), km, a, b])
    for (road, land), segs in sorted(groups.items()):
        segs.sort()
        lines = _chains(segs)
        geom = {"type": "MultiLineString", "coordinates": [[list(p) for p in ln] for ln in lines]}
        infra = "autostrada federale" if road.startswith("A") else "strada federale" if road.startswith("B") else "strada"
        props = {"road": road, "land": LAND.get(land, land), "infrastructure": infra, "segment_count": len(segs),
                 "length_km": round(sum(s[3] for s in segs), 1), "table_version": version, "valid_from": valid_from,
                 "operator": "Toll Collect GmbH", "vehicles": "veicoli > 3,5 t",
                 "segments": [[s[0], s[1], s[2], s[3]] for s in segs]}
        nid = f"{road}:{land}"
        yield NormalizedRecord(
            source_id=meta["source_id"], native_id=nid, native_version=content_version([props, geom]), kind="object",
            type="toll.road_section", label=f"{road} · {LAND.get(land, land)} — tratte a pedaggio (camion)",
            identifiers=[("balm_road_section", nid)], properties=props, geometry=geom,
            assertions=[Assertion("relation", "located_in", Target("place.country", scheme="iso3166a2", value="DE")),
                        Assertion("relation", "subject_to", Target("toll.scheme", scheme="toll_scheme", value="DE-LKW-MAUT")),
                        Assertion("relation", "priced_by", Target("toll.tariff", scheme="toll_tariff", value="DE-BFSTRMG-ANLAGE1"))],
            status="reviewed", method="asserted", raw_locator=f"Mauttabelle v{version} {road} {land}",
            text=f"{road} {LAND.get(land, land)} Maut")
