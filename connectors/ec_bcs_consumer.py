"""European Commission (DG ECFIN) Business and Consumer Surveys — consumer survey, monthly answer shares (NSA).
Licence: CC BY 4.0 (Commission legal notice). "Source: European Commission, DG ECFIN BCS".
From the published shares of each answer (++, +, =, −, −−, don't know) this connector computes, month by month, the
BALANCE exactly as the Commission defines it — B = (PP + ½P) − (½M + MM), "don't know" kept in the denominator —
and its sampling standard error at the Commission's nominal sample size (BCS User Guide, Table 2.1). Balances are not
percentages of people, and these are not seasonally adjusted: changes are compared with the SAME MONTH OF THE
PREVIOUS YEAR. The download link changes every month: the official page is read first (backfill rounds)."""

import io
import math
import re
import zipfile

import openpyxl

from connectors.obs_common import NAMES_IT, iso, record
from nexum.core.scheduler import FetchRequest

VERSION = "1.0.0"
PAGE = ("https://economy-finance.ec.europa.eu/economic-forecast-and-surveys/business-and-consumer-surveys/"
        "download-business-and-consumer-survey-data/time-series_en")
MONTHS_KEPT = 60
# nominal (effective target) consumer sample per country, BCS User Guide Table 2.1
N = {"BE": 1850, "BG": 1010, "CZ": 1000, "DK": 1100, "DE": 2020, "EE": 800, "IE": 1000, "EL": 1500, "ES": 2020, "FR": 1670,
     "HR": 1000, "IT": 2000, "CY": 600, "LV": 1000, "LT": 1110, "LU": 510, "HU": 1000, "MT": 1050, "NL": 1140, "AT": 1500,
     "PL": 1000, "PT": 1300, "RO": 1000, "SI": 940, "SK": 1200, "FI": 990, "SE": 1500, "ME": 1000, "MK": 1000, "AL": 1200,
     "RS": 1020, "TR": 3930}
TOPIC = {2: "famiglie", 3: "economia", 4: "economia", 5: "prezzi", 6: "prezzi", 7: "lavoro"}
QUESTIONS = {
    2: ("household.finance.expect12", "Situazione finanziaria della famiglia nei prossimi 12 mesi",
        "How do you expect the financial position of your household to change over the next 12 months?"),
    3: ("economy.past12", "Situazione economica generale negli ultimi 12 mesi",
        "How do you think the general economic situation in the country has changed over the past 12 months?"),
    4: ("economy.expect12.balance", "Situazione economica generale nei prossimi 12 mesi",
        "How do you expect the general economic situation in this country to develop over the next 12 months?"),
    5: ("prices.past12", "Prezzi negli ultimi 12 mesi (percezione)",
        "How do you think that consumer prices have developed over the last 12 months?"),
    6: ("prices.expect12", "Prezzi nei prossimi 12 mesi (attese)",
        "By comparison with the past 12 months, how do you expect that consumer prices will develop in the next 12 months?"),
    7: ("unemployment.expect12", "Disoccupazione nei prossimi 12 mesi (attese)",
        "How do you expect the number of people unemployed in this country to change over the next 12 months?"),
}
W = {"PP": 1.0, "P": 0.5, "E": 0.0, "M": -0.5, "MM": -1.0, "N": 0.0}


def describe():
    return {"connector_version": VERSION, "produces": ["observation.opinion_series"]}


def plan(mode, state, source, today=None):
    reqs = [FetchRequest(PAGE, "bcs_time_series_page.html")]
    if (state or {}).get("shares"):
        reqs.append(FetchRequest(state["shares"], "bcs_datashares_nsa_m.zip"))
    return reqs


def page_info(data):
    m = re.search(rb'href="(https://ec\.europa\.eu/economy_finance/db_indicators/surveys/documents/series/nace2_ecfin_\d{4}/datashares_nsa_m_nace2\.zip)"', data or b"")
    return {"shares": m.group(1).decode()} if m else None


def next_state(state, request, result, today=None):
    st = dict(state or {})
    if result and result.get("shares"):
        st["shares"] = result["shares"]
    return st


def parse(data: bytes, meta: dict):
    if data[:2] != b"PK":
        return                                        # the HTML page: only its link is used (page_info)
    z = zipfile.ZipFile(io.BytesIO(data))
    name = next((n for n in z.namelist() if n.startswith("datashares_consumers_nsa_m")), None)
    if not name:
        return
    wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True, data_only=True)
    cells = {}                                        # (country, q) -> {month: {cat: share}}
    for sh in wb.sheetnames:
        if not sh.startswith("CONSUMERS"):
            continue
        rows = wb[sh].iter_rows(values_only=True)
        hdr = next(rows)
        cols = []
        for j, h in enumerate(hdr):
            m = re.match(r"CONS\.([A-Z]{2})\.TOT\.(\d+)\.(PP|P|E|M|MM|N)\.M$", str(h or ""))
            if m and int(m.group(2)) in QUESTIONS and m.group(1) in N:
                cols.append((j, m.group(1), int(m.group(2)), m.group(3)))
        if not cols:
            continue
        for r in rows:
            d = r[0]
            if not hasattr(d, "year"):
                continue
            month = f"{d.year:04d}-{d.month:02d}"
            for j, c, q, cat in cols:
                v = r[j]
                if isinstance(v, (int, float)):
                    cells.setdefault((c, q), {}).setdefault(month, {})[cat] = v
    for (c, q), months in sorted(cells.items()):
        code, label, question = QUESTIONS[q]
        n = N[c]
        pts = []
        for month in sorted(months)[-MONTHS_KEPT:]:
            sh = months[month]
            tot = sum(sh.get(k, 0) for k in W)
            if tot < 90:                               # incomplete month (missing answer categories)
                continue
            p = {k: sh.get(k, 0) / tot for k in W}
            mean = sum(W[k] * p[k] for k in W)
            var = sum(W[k] ** 2 * p[k] for k in W) - mean ** 2
            y, mo = int(month[:4]), int(month[5:])
            start, end = f"{month}-01", f"{month}-{[31, 29 if y % 4 == 0 else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][mo - 1]:02d}"
            pts.append([start, end, round(100 * mean, 1), n, round(100 * math.sqrt(max(var, 0) / n), 2)])
        if len(pts) < 3:
            continue
        i2 = iso(c)
        name = NAMES_IT.get(i2, i2)
        yield record(meta, f"bcs:{code}:{i2}", "observation.opinion_series", f"{label} · {name} (Commissione UE, saldo)", i2,
                     {"indicator": code, "indicator_label": label, "topic": TOPIC[q], "question": question,
                      "answer": "saldo: risposte positive meno negative (−100…+100)", "statistic": "balance", "unit": "saldo",
                      "population": "consumatori (campione nazionale, obiettivo effettivo %d interviste)" % n,
                      "method": "indagine armonizzata UE; quote di risposta pubblicate dalla Commissione; errore standard calcolato al campione nominale",
                      "probability_sample": True, "seasonal": "non destagionalizzato",
                      "compare": "same_month_previous_year", "dataset": "Commissione europea, DG ECFIN — Business and Consumer Surveys",
                      "comparable_series_id": f"bcs:{code}:{i2}", "series": pts, "latest_value": pts[-1][2], "latest_end": pts[-1][1]},
                     text=f"{label} {name} consumatori Commissione europea")
